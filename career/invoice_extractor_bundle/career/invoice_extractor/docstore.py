"""Download invoice PDFs from bosuka DocStore (Selenium + basic auth).

Working pattern (matches the Aurora/Selenium script):
  1. Open the document page in Chrome with credentials in the URL
  2. Scrape <a href> links that contain both "pdf" and "filename"
  3. Download the PDF bytes with requests + HTTP basic auth

Credentials: DOCSTORE_USERNAME / DOCSTORE_PASSWORD (never hardcode).

Excel input: each row supplies DocRef (LREF) and Country, e.g. AT + 00011841
→ http://.../store/AT%20Docstore/document/?L[LREF]=00011841
"""
from __future__ import annotations

import os
import re
from abc import ABC, abstractmethod
from pathlib import Path
from urllib.parse import quote

import requests

# Header aliases accepted when reading an Excel job list
_DOCREF_HEADERS = {
    "docref",
    "doc ref",
    "doc_ref",
    "allrows.docref",
    "lref",
    "l[lref]",
    "document ref",
    "document_ref",
}
_COUNTRY_HEADERS = {
    "country",
    "ctry",
    "co",
    "docstore country",
    "docstore_country",
}


class DocstoreClient(ABC):
    """Fetch invoice PDFs from a document store."""

    @abstractmethod
    def list_pdfs(self) -> list[str]:
        ...

    @abstractmethod
    def download(self, doc_id: str, dest_dir: Path) -> Path:
        ...

    def close(self) -> None:
        return None


class LocalFolderDocstore(DocstoreClient):
    """Read PDFs already on disk (no login)."""

    def __init__(self, folder: Path) -> None:
        self.folder = Path(folder)
        if not self.folder.is_dir():
            raise NotADirectoryError(self.folder)

    def list_pdfs(self) -> list[str]:
        return sorted(p.name for p in self.folder.glob("*.pdf"))

    def download(self, doc_id: str, dest_dir: Path) -> Path:
        src = self.folder / doc_id
        if not src.is_file():
            raise FileNotFoundError(src)
        return src


def _norm_header(value: object) -> str:
    text = str(value or "").strip().lower()
    text = text.replace("_", " ")
    return re.sub(r"\s+", " ", text)


def _cell_to_docref(value: object) -> str:
    """Preserve leading zeros when Excel stored the ref as text."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return ""
    if isinstance(value, float) and value == int(value):
        return str(int(value))
    if isinstance(value, int):
        return str(value)
    return str(value).strip()


def _cell_to_country(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip().upper()


def read_jobs_from_excel(
    path: Path,
    sheet: str | None = None,
    *,
    docref_col: str | None = None,
    country_col: str | None = None,
) -> list[tuple[str, str]]:
    """Read (country, docref) pairs from an Excel sheet.

    Looks for columns named like DocRef / AllRows.DOCREF / LREF and Country.
    Blank rows and duplicates are skipped (first occurrence kept).
    """
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ImportError("openpyxl is required to read Excel job lists") from exc

    path = Path(path)
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet:
            if sheet not in wb.sheetnames:
                raise ValueError(
                    f"Sheet {sheet!r} not in {path.name}. Available: {wb.sheetnames}"
                )
            ws = wb[sheet]
        else:
            ws = _find_sheet_with_docref(wb)
            if ws is None:
                raise ValueError(
                    f"No sheet in {path.name} has DocRef/Country columns. "
                    f"Sheets: {wb.sheetnames}"
                )

        rows_iter = ws.iter_rows(values_only=True)
        header_row = next(rows_iter, None)
        if not header_row:
            raise ValueError(f"Sheet {ws.title!r} is empty")

        headers = [_norm_header(h) for h in header_row]
        # Also keep dotted forms (allrows.docref) after normalizing spaces
        headers_compact = [h.replace(" ", "") for h in headers]

        def find_col(aliases: set[str], explicit: str | None) -> int:
            if explicit:
                want = _norm_header(explicit)
                for i, h in enumerate(headers):
                    if h == want or h.replace(" ", "") == want.replace(" ", ""):
                        return i
                raise ValueError(f"Column {explicit!r} not found in {ws.title!r}")
            for i, h in enumerate(headers):
                if h in aliases or h.replace(" ", "") in {a.replace(" ", "") for a in aliases}:
                    return i
            for i, h in enumerate(headers_compact):
                if h in {a.replace(" ", "") for a in aliases}:
                    return i
            raise ValueError(
                f"Could not find DocRef/Country column in {ws.title!r}. "
                f"Headers: {[str(x) for x in header_row if x]}"
            )

        doc_i = find_col(_DOCREF_HEADERS, docref_col)
        co_i = find_col(_COUNTRY_HEADERS, country_col)

        jobs: list[tuple[str, str]] = []
        seen: set[tuple[str, str]] = set()
        for row in rows_iter:
            if row is None:
                continue
            docref = _cell_to_docref(row[doc_i] if doc_i < len(row) else None)
            country = _cell_to_country(row[co_i] if co_i < len(row) else None)
            if not docref or not country:
                continue
            key = (country, docref)
            if key in seen:
                continue
            seen.add(key)
            jobs.append(key)
        return jobs
    finally:
        wb.close()


def _find_sheet_with_docref(wb) -> object | None:
    for name in wb.sheetnames:
        ws = wb[name]
        row1 = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
        if not row1:
            continue
        headers = {_norm_header(h) for h in row1 if h is not None}
        headers |= {h.replace(" ", "") for h in headers}
        has_doc = bool(headers & _DOCREF_HEADERS) or bool(
            headers & {a.replace(" ", "") for a in _DOCREF_HEADERS}
        )
        has_co = bool(headers & _COUNTRY_HEADERS)
        if has_doc and has_co:
            return ws
    return None


class BosukaDocstore(DocstoreClient):
    """New Balance DocStore on bosuka1 via Selenium link discovery.

    Document URL shape:
      http://USER:PASS@bosuka1.newbalance.com:6400/docStore/store/AT%20Docstore/document/?L[LREF]=00011841

    Supports a single country + LREF list, or mixed (country, docref) jobs from Excel.
    """

    def __init__(
        self,
        username: str,
        password: str,
        lrefs: list[str] | None = None,
        country: str = "NG",
        jobs: list[tuple[str, str]] | None = None,
        base_url: str = "http://bosuka1.newbalance.com:6400",
        headless: bool = True,
    ) -> None:
        if not username or not password:
            raise ValueError(
                "DocStore credentials required. Set DOCSTORE_USERNAME and "
                "DOCSTORE_PASSWORD in career/.env"
            )
        self.username = username
        self.password = password
        self.base_url = base_url.rstrip("/")
        self.headless = headless
        self._driver = None
        # key = "AT:00011841" -> pdf url
        self._pdf_urls: dict[str, str] = {}

        seen: set[str] = set()
        self.jobs: list[tuple[str, str]] = []
        if jobs:
            for raw_country, raw_lref in jobs:
                country_code = str(raw_country).strip().upper()
                lref = str(raw_lref).strip()
                if not country_code or not lref:
                    continue
                key = f"{country_code}:{lref}"
                if key in seen:
                    continue
                seen.add(key)
                self.jobs.append((country_code, lref))
        else:
            default_country = country.strip().upper()
            for raw in lrefs or []:
                lref = str(raw).strip()
                if not lref:
                    continue
                key = f"{default_country}:{lref}"
                if key in seen:
                    continue
                seen.add(key)
                self.jobs.append((default_country, lref))

    @staticmethod
    def job_id(country: str, lref: str) -> str:
        return f"{country.strip().upper()}:{str(lref).strip()}"

    @staticmethod
    def parse_job_id(doc_id: str) -> tuple[str, str]:
        if ":" not in doc_id:
            raise ValueError(
                f"Expected job id like 'AT:00011841', got {doc_id!r}"
            )
        country, lref = doc_id.split(":", 1)
        return country.strip().upper(), lref.strip()

    def store_name(self, country: str) -> str:
        return f"{country.strip().upper()} Docstore"

    def document_page_url(
        self, lref: str, country: str, *, with_auth: bool = True
    ) -> str:
        store = quote(self.store_name(country), safe="")
        path = (
            f"{self.base_url}/docStore/store/{store}/document/"
            f"?L[LREF]={quote(str(lref), safe='')}"
        )
        if not with_auth:
            return path
        cred = f"{quote(self.username, safe='')}:{quote(self.password, safe='')}"
        return path.replace("://", f"://{cred}@", 1)

    def list_pdfs(self) -> list[str]:
        if not self.jobs:
            raise ValueError(
                "No DocStore jobs. Pass --from-excel, set DOCSTORE_LREFS, "
                "or use --lref / --country"
            )
        return [self.job_id(c, l) for c, l in self.jobs]

    def download(self, doc_id: str, dest_dir: Path) -> Path:
        dest_dir = Path(dest_dir)
        dest_dir.mkdir(parents=True, exist_ok=True)
        country, lref = self.parse_job_id(doc_id)

        pdf_url = self._pdf_urls.get(doc_id)
        if pdf_url is None:
            self._discover_pdf_urls([(country, lref)])
            pdf_url = self._pdf_urls.get(doc_id)
        if not pdf_url:
            raise RuntimeError(
                f"No PDF link found for {country} LREF={lref}. "
                f"Page: {self.document_page_url(lref, country, with_auth=False)}"
            )

        resp = requests.get(
            pdf_url,
            auth=(self.username, self.password),
            timeout=120,
        )
        resp.raise_for_status()
        if resp.content[:4] != b"%PDF":
            raise RuntimeError(
                f"Download for {country}_{lref} did not return a PDF "
                f"(content-type={resp.headers.get('Content-Type')!r})"
            )

        out = dest_dir / f"{country}_{lref}.pdf"
        out.write_bytes(resp.content)
        return out

    def prefetch(self) -> None:
        """Open one Chrome session and resolve PDF links for all jobs."""
        self._discover_pdf_urls(self.jobs)

    def _ensure_driver(self):
        if self._driver is not None:
            return self._driver
        try:
            from selenium import webdriver
            from selenium.webdriver.chrome.options import Options
        except ImportError as exc:
            raise ImportError(
                "selenium is required for DocStore downloads. "
                "pip install selenium"
            ) from exc

        options = Options()
        if self.headless:
            options.add_argument("--headless=new")
        options.add_argument("--disable-gpu")
        options.add_argument("--no-sandbox")
        self._driver = webdriver.Chrome(options=options)
        return self._driver

    def _discover_pdf_urls(self, jobs: list[tuple[str, str]]) -> None:
        from selenium.webdriver.common.by import By

        driver = self._ensure_driver()
        for country, lref in jobs:
            key = self.job_id(country, lref)
            page_url = self.document_page_url(lref, country, with_auth=True)
            driver.get(page_url)
            pdf_href: str | None = None
            for link in driver.find_elements(By.XPATH, "//a[@href]"):
                href = link.get_attribute("href") or ""
                if "pdf" in href.lower() and "filename" in href.lower():
                    pdf_href = href
            if pdf_href:
                self._pdf_urls[key] = pdf_href
            else:
                print(f"Warning: no PDF link for {country} LREF={lref}")

    def close(self) -> None:
        if self._driver is not None:
            self._driver.quit()
            self._driver = None


def default_invoice_folder() -> Path:
    """Prefer the NB OneDrive Invoice folder when present."""
    configured = os.getenv("INVOICE_DOWNLOAD_DIR") or os.getenv("DOCSTORE_LOCAL_FOLDER")
    if configured:
        return Path(configured)

    candidates = [
        Path(r"C:\Users\shjiang\OneDrive - New Balance Athletics, Inc\Documents\Invoice"),
        Path.home()
        / "OneDrive - New Balance Athletics, Inc"
        / "Documents"
        / "Invoice",
        Path("career/invoice_pdfs"),
    ]
    for path in candidates:
        if path.is_dir():
            return path
    return candidates[0]


def build_docstore_from_env(
    extra_lrefs: list[str] | None = None,
    jobs: list[tuple[str, str]] | None = None,
) -> DocstoreClient:
    """Build client from environment variables and/or Excel jobs."""
    mode = os.getenv("DOCSTORE_MODE", "local").lower()
    if mode == "local" and not jobs and not extra_lrefs:
        return LocalFolderDocstore(default_invoice_folder())

    if mode in {"nf", "nfdocstore", "bosuka", "docstore"} or jobs or extra_lrefs:
        lrefs_env = os.getenv("DOCSTORE_LREFS", "")
        lrefs = [x.strip() for x in lrefs_env.split(",") if x.strip()]
        if extra_lrefs:
            lrefs.extend(extra_lrefs)
        headless = os.getenv("DOCSTORE_HEADLESS", "1").strip().lower() not in {
            "0",
            "false",
            "no",
        }
        return BosukaDocstore(
            username=os.getenv("DOCSTORE_USERNAME", ""),
            password=os.getenv("DOCSTORE_PASSWORD", ""),
            lrefs=None if jobs else lrefs,
            country=os.getenv("DOCSTORE_COUNTRY", "NG"),
            jobs=jobs,
            base_url=os.getenv(
                "DOCSTORE_BASE_URL",
                "http://bosuka1.newbalance.com:6400",
            ),
            headless=headless,
        )

    raise ValueError(
        f"Unknown DOCSTORE_MODE={mode!r} (use local or bosuka/nf)"
    )
